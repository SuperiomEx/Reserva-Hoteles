from rest_framework import viewsets, filters, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.views import APIView
from django_filters.rest_framework import DjangoFilterBackend
from django.db.models import Q, Count, Sum
from django.http import HttpResponse
from django.contrib.auth.models import User
from django.contrib.auth import authenticate
from django.utils import timezone
import csv
from datetime import datetime, timedelta
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from io import BytesIO

from .models import TipoHabitacion, Huesped, Habitacion, Reservacion
from .serializers import (
    TipoHabitacionSerializer, 
    HuespedSerializer, 
    HabitacionSerializer, 
    ReservacionSerializer,
    ReservacionCreateSerializer,
    ReservacionFilterSerializer
)

class TipoHabitacionViewSet(viewsets.ModelViewSet):
    """ViewSet para gestión de tipos de habitación"""
    queryset = TipoHabitacion.objects.all()
    serializer_class = TipoHabitacionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['nombre', 'descripcion']
    ordering_fields = ['nombre', 'precio_base', 'capacidad_personas']
    ordering = ['nombre']

class HuespedViewSet(viewsets.ModelViewSet):
    """ViewSet para gestión de huéspedes"""
    queryset = Huesped.objects.all()
    serializer_class = HuespedSerializer
    permission_classes = [AllowAny]  # Temporalmente permitir sin auth
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['sexo', 'activo', 'nacionalidad']
    search_fields = ['nombres', 'apellidos', 'numero_identidad', 'correo_electronico']
    ordering_fields = ['nombres', 'apellidos', 'created_at']
    ordering = ['apellidos', 'nombres']
    
    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)
    
    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)

class HabitacionViewSet(viewsets.ModelViewSet):
    """ViewSet para gestión de habitaciones"""
    queryset = Habitacion.objects.select_related('tipo_habitacion').all()
    serializer_class = HabitacionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['tipo_habitacion', 'estado']
    search_fields = ['numero_habitacion', 'descripcion']
    ordering_fields = ['numero_habitacion', 'created_at']
    ordering = ['numero_habitacion']
    
    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)
    
    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)
    
    @action(detail=False, methods=['get'])
    def disponibles(self, request):
        """Obtener habitaciones disponibles"""
        habitaciones = self.queryset.filter(estado='DISPONIBLE')
        serializer = self.get_serializer(habitaciones, many=True)
        return Response(serializer.data)

class ReservacionViewSet(viewsets.ModelViewSet):
    """ViewSet para gestión de reservaciones"""
    queryset = Reservacion.objects.select_related('huesped', 'habitacion').all()
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['estado', 'habitacion', 'huesped']
    search_fields = ['numero_confirmacion', 'huesped__nombres', 'huesped__apellidos', 
                    'habitacion__numero_habitacion']
    ordering_fields = ['fecha_llegada', 'fecha_salida', 'created_at', 'precio'
    ]
    ordering = ['-fecha_llegada']
    
    def get_serializer_class(self):
        if self.action in ['create', 'update', 'partial_update']:
            return ReservacionCreateSerializer
        return ReservacionSerializer
    
    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)
    
    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)
    
    def get_queryset(self):
        queryset = super().get_queryset()
        
        # Filtros personalizados
        huesped_id = self.request.query_params.get('huesped_id')
        habitacion_numero = self.request.query_params.get('numero_habitacion')
        activa = self.request.query_params.get('activa')
        
        if huesped_id:
            queryset = queryset.filter(huesped_id=huesped_id)
        
        if habitacion_numero:
            queryset = queryset.filter(habitacion__numero_habitacion__icontains=habitacion_numero)
        
        if activa is not None:
            if activa.lower() == 'true':
                queryset = queryset.filter(estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO'])
            else:
                queryset = queryset.exclude(estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO'])
        
        return queryset
    
    @action(detail=False, methods=['get'])
    def activas(self, request):
        """Obtener reservaciones activas"""
        reservaciones = self.get_queryset().filter(
            estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO']
        )
        
        page = self.paginate_queryset(reservaciones)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        
        serializer = self.get_serializer(reservaciones, many=True)
        return Response(serializer.data)
    
    @action(detail=False, methods=['get'])
    def export_activas_csv(self, request):
        """Exportar reservaciones activas en formato CSV"""
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="reservaciones_activas_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv"'
        
        writer = csv.writer(response)
        writer.writerow([
            'Número Confirmación', 'Huésped', 'Habitación', 'Fecha Llegada', 
            'Fecha Salida', 'Precio', 'Estado', 'Método Pago'
        ])
        
        reservaciones = self.get_queryset().filter(
            estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO']
        )
        
        for reserva in reservaciones:
            writer.writerow([
                reserva.numero_confirmacion,
                reserva.huesped.nombre_completo,
                reserva.habitacion.numero_habitacion,
                reserva.fecha_llegada,
                reserva.fecha_salida,
                reserva.precio,
                reserva.get_estado_display(),
                reserva.get_metodo_pago_display() if reserva.metodo_pago else ''
            ])
        
        return response
    
    @action(detail=False, methods=['post'])
    def buscar(self, request):
        """Búsqueda avanzada de reservaciones"""
        serializer = ReservacionFilterSerializer(data=request.data)
        if serializer.is_valid():
            queryset = self.get_queryset()
            
            # Aplicar filtros
            if serializer.validated_data.get('huesped_id'):
                queryset = queryset.filter(huesped_id=serializer.validated_data['huesped_id'])
            
            if serializer.validated_data.get('habitacion_id'):
                queryset = queryset.filter(habitacion_id=serializer.validated_data['habitacion_id'])
            
            if serializer.validated_data.get('numero_habitacion'):
                queryset = queryset.filter(
                    habitacion__numero_habitacion__icontains=serializer.validated_data['numero_habitacion']
                )
            
            if serializer.validated_data.get('activa') is not None:
                if serializer.validated_data['activa']:
                    queryset = queryset.filter(estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO'])
                else:
                    queryset = queryset.exclude(estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO'])
            
            if serializer.validated_data.get('estado'):
                queryset = queryset.filter(estado=serializer.validated_data['estado'])
            
            if serializer.validated_data.get('fecha_llegada_desde'):
                queryset = queryset.filter(fecha_llegada__gte=serializer.validated_data['fecha_llegada_desde'])
            
            if serializer.validated_data.get('fecha_llegada_hasta'):
                queryset = queryset.filter(fecha_llegada__lte=serializer.validated_data['fecha_llegada_hasta'])
            
            page = self.paginate_queryset(queryset)
            if page is not None:
                result_serializer = self.get_serializer(page, many=True)
                return self.get_paginated_response(result_serializer.data)
            
            result_serializer = self.get_serializer(queryset, many=True)
            return Response(result_serializer.data)
        
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

# Vistas de autenticación
class RegisterView(APIView):
    """Vista para registro de usuarios"""
    permission_classes = [AllowAny]
    
    def post(self, request):
        username = request.data.get('username')
        email = request.data.get('email')
        password = request.data.get('password')
        first_name = request.data.get('first_name', '')
        last_name = request.data.get('last_name', '')
        
        if not username or not email or not password:
            return Response({
                'error': 'Username, email y password son requeridos'
            }, status=status.HTTP_400_BAD_REQUEST)
        
        if User.objects.filter(username=username).exists():
            return Response({
                'error': 'El username ya existe'
            }, status=status.HTTP_400_BAD_REQUEST)
        
        if User.objects.filter(email=email).exists():
            return Response({
                'error': 'El email ya está registrado'
            }, status=status.HTTP_400_BAD_REQUEST)
        
        user = User.objects.create_user(
            username=username,
            email=email,
            password=password,
            first_name=first_name,
            last_name=last_name
        )
        
        # Crear token
        from rest_framework.authtoken.models import Token
        token, created = Token.objects.get_or_create(user=user)
        
        return Response({
            'access': token.key,
            'refresh': token.key,  # Para compatibilidad
            'user': {
                'id': user.id,
                'username': user.username,
                'email': user.email,
                'first_name': user.first_name,
                'last_name': user.last_name,
                'is_active': user.is_active,
                'date_joined': user.date_joined.isoformat()
            }
        }, status=status.HTTP_201_CREATED)

class UserProfileView(APIView):
    """Vista para obtener el perfil del usuario"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return Response({
            'id': request.user.id,
            'username': request.user.username,
            'email': request.user.email,
            'first_name': request.user.first_name,
            'last_name': request.user.last_name,
            'is_staff': request.user.is_staff
        })

# Vista para estadísticas del dashboard
class DashboardStatsView(APIView):
    """Vista para obtener estadísticas del dashboard"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        hoy = timezone.now().date()
        
        # Estadísticas básicas
        total_habitaciones = Habitacion.objects.count()
        habitaciones_disponibles = Habitacion.objects.filter(estado='DISPONIBLE').count()
        habitaciones_ocupadas = Habitacion.objects.filter(estado='OCUPADA').count()
        
        total_huespedes = Huesped.objects.filter(activo=True).count()
        
        # Reservaciones
        reservaciones_activas = Reservacion.objects.filter(
            estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO']
        ).count()
        
        reservaciones_hoy = Reservacion.objects.filter(
            fecha_llegada=hoy,
            estado__in=['PENDIENTE', 'CONFIRMADA']
        ).count()
        
        check_outs_hoy = Reservacion.objects.filter(
            fecha_salida=hoy,
            estado='EN_CURSO'
        ).count()
        
        # Ingresos del mes actual
        inicio_mes = hoy.replace(day=1)
        ingresos_mes = Reservacion.objects.filter(
            fecha_llegada__gte=inicio_mes,
            estado__in=['CONFIRMADA', 'EN_CURSO', 'COMPLETADA']
        ).aggregate(total=Sum('precio'))['total'] or 0
        
        # Ocupación por tipo de habitación
        ocupacion_tipos = TipoHabitacion.objects.annotate(
            total_habitaciones=Count('habitacion'),
            ocupadas=Count('habitacion', filter=Q(habitacion__estado='OCUPADA'))
        ).values('nombre', 'total_habitaciones', 'ocupadas')
        
        return Response({
            'habitaciones': {
                'total': total_habitaciones,
                'disponibles': habitaciones_disponibles,
                'ocupadas': habitaciones_ocupadas,
                'porcentaje_ocupacion': round((habitaciones_ocupadas / total_habitaciones * 100) if total_habitaciones > 0 else 0, 2)
            },
            'huespedes': {
                'total_activos': total_huespedes
            },
            'reservaciones': {
                'activas': reservaciones_activas,
                'check_ins_hoy': reservaciones_hoy,
                'check_outs_hoy': check_outs_hoy
            },
            'ingresos': {
                'mes_actual': float(ingresos_mes)
            },
            'ocupacion_por_tipo': list(ocupacion_tipos)
        })

# Vista para exportar PDF
class ExportReservacionesPDFView(APIView):
    """Vista para exportar reservaciones activas en PDF"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        # Obtener reservaciones activas
        reservaciones = Reservacion.objects.filter(
            estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO']
        ).select_related('huesped', 'habitacion').order_by('fecha_llegada')
        
        # Crear el PDF
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []
        
        # Estilos
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            alignment=1  # Centrado
        )
        
        # Título
        title = Paragraph("Listado de Reservaciones Activas", title_style)
        elements.append(title)
        elements.append(Spacer(1, 20))
        
        # Fecha de generación
        fecha_generacion = Paragraph(
            f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            styles['Normal']
        )
        elements.append(fecha_generacion)
        elements.append(Spacer(1, 20))
        
        # Datos de la tabla
        data = [['Confirmación', 'Huésped', 'Habitación', 'Llegada', 'Salida', 'Precio', 'Estado']]
        
        for reserva in reservaciones:
            data.append([
                reserva.numero_confirmacion,
                reserva.huesped.nombre_completo,
                reserva.habitacion.numero_habitacion,
                reserva.fecha_llegada.strftime('%d/%m/%Y'),
                reserva.fecha_salida.strftime('%d/%m/%Y'),
                f"${reserva.precio:,.2f}",
                reserva.get_estado_display()
            ])
        
        # Crear tabla
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(table)
        
        # Total de reservaciones
        elements.append(Spacer(1, 20))
        total_reservaciones = Paragraph(
            f"Total de reservaciones activas: {reservaciones.count()}",
            styles['Normal']
        )
        elements.append(total_reservaciones)
        
        # Construir PDF
        doc.build(elements)
        buffer.seek(0)
        
        # Respuesta HTTP mejorada para Chrome
        response = HttpResponse(buffer.read(), content_type='application/pdf')
        
        # Headers adicionales para Chrome
        filename = f"reservaciones_activas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        response['Content-Length'] = len(buffer.getvalue())
        response['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response['Pragma'] = 'no-cache'
        response['Expires'] = '0'
        response['X-Content-Type-Options'] = 'nosniff'
        
        return response

# Vista separada para exportar Excel
class ExportReservacionesExcelView(APIView):
    """Vista para exportar reservaciones activas en Excel"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Reservaciones Activas"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        center_alignment = Alignment(horizontal="center")
        
        # Encabezados
        headers = [
            'Número Confirmación', 'Huésped', 'Habitación', 'Fecha Llegada', 
            'Fecha Salida', 'Precio', 'Estado', 'Método Pago'
        ]
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        # Datos
        reservaciones = Reservacion.objects.filter(
            estado__in=['PENDIENTE', 'CONFIRMADA', 'EN_CURSO']
        ).select_related('huesped', 'habitacion')
        
        for row_num, reserva in enumerate(reservaciones, 2):
            ws.cell(row=row_num, column=1, value=reserva.numero_confirmacion)
            ws.cell(row=row_num, column=2, value=reserva.huesped.nombre_completo)
            ws.cell(row=row_num, column=3, value=reserva.habitacion.numero_habitacion)
            ws.cell(row=row_num, column=4, value=reserva.fecha_llegada)
            ws.cell(row=row_num, column=5, value=reserva.fecha_salida)
            ws.cell(row=row_num, column=6, value=float(reserva.precio))
            ws.cell(row=row_num, column=7, value=reserva.get_estado_display())
            ws.cell(row=row_num, column=8, value=reserva.get_metodo_pago_display() if reserva.metodo_pago else '')
        
        # Ajustar ancho de columnas
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Guardar en memoria
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Respuesta HTTP
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="reservaciones_activas_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        
        return response

class ReportesViewSet(APIView):
    """Vista para generar reportes específicos"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        tipo_reporte = request.GET.get('tipo', 'ocupacion')
        fecha_inicio = request.GET.get('fecha_inicio')
        fecha_fin = request.GET.get('fecha_fin')
        
        if not fecha_inicio or not fecha_fin:
            return Response({'error': 'Fechas requeridas'}, status=400)
            
        try:
            fecha_inicio = datetime.strptime(fecha_inicio, '%Y-%m-%d').date()
            fecha_fin = datetime.strptime(fecha_fin, '%Y-%m-%d').date()
        except ValueError:
            return Response({'error': 'Formato de fecha inválido'}, status=400)
            
        if tipo_reporte == 'ocupacion':
            return self.reporte_ocupacion(fecha_inicio, fecha_fin)
        elif tipo_reporte == 'ingresos':
            return self.reporte_ingresos(fecha_inicio, fecha_fin)
        elif tipo_reporte == 'huespedes':
            return self.reporte_huespedes(fecha_inicio, fecha_fin)
        elif tipo_reporte == 'habitaciones':
            return self.reporte_habitaciones()
        else:
            return Response({'error': 'Tipo de reporte no válido'}, status=400)
    
    def reporte_ocupacion(self, fecha_inicio, fecha_fin):
        """Reporte de ocupación hotelera"""
        total_habitaciones = Habitacion.objects.count()
        
        # Reservaciones en el período
        reservaciones_periodo = Reservacion.objects.filter(
            fecha_llegada__gte=fecha_inicio,
            fecha_salida__lte=fecha_fin,
            estado__in=['CONFIRMADA', 'EN_CURSO', 'COMPLETADA']
        )
        
        habitaciones_ocupadas = reservaciones_periodo.values('habitacion').distinct().count()
        habitaciones_disponibles = total_habitaciones - habitaciones_ocupadas
        tasa_ocupacion = (habitaciones_ocupadas / total_habitaciones * 100) if total_habitaciones > 0 else 0
        
        # Ocupación por tipo de habitación
        ocupacion_por_tipo = TipoHabitacion.objects.annotate(
            total_hab=Count('habitacion'),
            ocupadas_periodo=Count(
                'habitacion__reservacion',
                filter=Q(
                    habitacion__reservacion__fecha_llegada__gte=fecha_inicio,
                    habitacion__reservacion__fecha_salida__lte=fecha_fin,
                    habitacion__reservacion__estado__in=['CONFIRMADA', 'EN_CURSO', 'COMPLETADA']
                )
            )
        ).values('nombre', 'total_hab', 'ocupadas_periodo')
        
        return Response({
            'habitaciones_ocupadas': habitaciones_ocupadas,
            'habitaciones_disponibles': habitaciones_disponibles,
            'tasa_ocupacion': round(tasa_ocupacion, 2),
            'total_habitaciones': total_habitaciones,
            'ocupacion_por_tipo': list(ocupacion_por_tipo),
            'periodo': {
                'inicio': fecha_inicio,
                'fin': fecha_fin
            }
        })
    
    def reporte_ingresos(self, fecha_inicio, fecha_fin):
        """Reporte de ingresos del hotel"""
        reservaciones_periodo = Reservacion.objects.filter(
            fecha_llegada__gte=fecha_inicio,
            fecha_salida__lte=fecha_fin,
            estado__in=['CONFIRMADA', 'EN_CURSO', 'COMPLETADA']
        )
        
        ingresos_periodo = reservaciones_periodo.aggregate(
            total=Sum('precio')
        )['total'] or 0
        
        reservaciones_confirmadas = reservaciones_periodo.count()
        promedio_reservacion = (ingresos_periodo / reservaciones_confirmadas) if reservaciones_confirmadas > 0 else 0
        
        # Ingresos por tipo de habitación
        ingresos_por_tipo = reservaciones_periodo.values(
            'habitacion__tipo__nombre'
        ).annotate(
            total_ingresos=Sum('precio'),
            cantidad_reservas=Count('id')
        ).order_by('-total_ingresos')
        
        # Ingresos por mes
        ingresos_mensuales = reservaciones_periodo.extra(
            select={'mes': "DATE_TRUNC('month', fecha_llegada)"}
        ).values('mes').annotate(
            ingresos=Sum('precio'),
            reservas=Count('id')
        ).order_by('mes')
        
        return Response({
            'ingresos_periodo': float(ingresos_periodo),
            'reservaciones_confirmadas': reservaciones_confirmadas,
            'promedio_reservacion': float(promedio_reservacion),
            'ingresos_por_tipo': list(ingresos_por_tipo),
            'ingresos_mensuales': list(ingresos_mensuales),
            'periodo': {
                'inicio': fecha_inicio,
                'fin': fecha_fin
            }
        })
    
    def reporte_huespedes(self, fecha_inicio, fecha_fin):
        """Reporte de huéspedes y actividad"""
        # Huéspedes activos en el período
        huespedes_periodo = Huesped.objects.filter(
            reservacion__fecha_llegada__gte=fecha_inicio,
            reservacion__fecha_salida__lte=fecha_fin
        ).distinct()
        
        total_huespedes = huespedes_periodo.count()
        
        # Nuevos huéspedes (primera reserva en el período)
        nuevos_huespedes = 0
        for huesped in huespedes_periodo:
            primera_reserva = huesped.reservacion_set.order_by('fecha_llegada').first()
            if primera_reserva and primera_reserva.fecha_llegada >= fecha_inicio:
                nuevos_huespedes += 1
        
        # Check-ins y check-outs del día actual
        hoy = timezone.now().date()
        checkins_hoy = Reservacion.objects.filter(
            fecha_llegada=hoy,
            estado__in=['PENDIENTE', 'CONFIRMADA']
        ).count()
        
        checkouts_hoy = Reservacion.objects.filter(
            fecha_salida=hoy,
            estado='EN_CURSO'
        ).count()
        
        # Huéspedes por nacionalidad/país (si tienes el campo)
        huespedes_por_pais = huespedes_periodo.values(
            'pais'
        ).annotate(
            cantidad=Count('id')
        ).order_by('-cantidad')[:10]
        
        return Response({
            'total_huespedes': total_huespedes,
            'nuevos_huespedes': nuevos_huespedes,
            'checkins_hoy': checkins_hoy,
            'checkouts_hoy': checkouts_hoy,
            'huespedes_por_pais': list(huespedes_por_pais),
            'periodo': {
                'inicio': fecha_inicio,
                'fin': fecha_fin
            }
        })
    
    def reporte_habitaciones(self):
        """Reporte del estado de las habitaciones"""
        total_habitaciones = Habitacion.objects.count()
        
        # Estado actual de habitaciones
        habitaciones_por_estado = Habitacion.objects.values(
            'estado'
        ).annotate(
            cantidad=Count('id')
        )
        
        # Habitaciones por tipo
        habitaciones_por_tipo = TipoHabitacion.objects.annotate(
            total_habitaciones=Count('habitacion'),
            disponibles=Count('habitacion', filter=Q(habitacion__estado='DISPONIBLE')),
            ocupadas=Count('habitacion', filter=Q(habitacion__estado='OCUPADA')),
            mantenimiento=Count('habitacion', filter=Q(habitacion__estado='MANTENIMIENTO'))
        ).values(
            'nombre', 'total_habitaciones', 'disponibles', 'ocupadas', 'mantenimiento'
        )
        
        return Response({
            'total_habitaciones': total_habitaciones,
            'habitaciones_disponibles': Habitacion.objects.filter(estado='DISPONIBLE').count(),
            'habitaciones_mantenimiento': Habitacion.objects.filter(estado='MANTENIMIENTO').count(),
            'habitaciones_limpieza': Habitacion.objects.filter(estado='LIMPIEZA').count(),
            'habitaciones_por_estado': list(habitaciones_por_estado),
            'habitaciones_por_tipo': list(habitaciones_por_tipo)
        })

class ExportReportsView(APIView):
    """Vista para exportar reportes en PDF y Excel"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request, format_type):
        tipo_reporte = request.GET.get('tipo', 'ocupacion')
        fecha_inicio = request.GET.get('fecha_inicio')
        fecha_fin = request.GET.get('fecha_fin')
        
        if format_type == 'pdf':
            return self.export_pdf(tipo_reporte, fecha_inicio, fecha_fin)
        elif format_type == 'excel':
            return self.export_excel(tipo_reporte, fecha_inicio, fecha_fin)
        else:
            return Response({'error': 'Formato no soportado'}, status=400)
    
    def export_pdf(self, tipo_reporte, fecha_inicio, fecha_fin):
        """Exportar reporte en PDF"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []
        styles = getSampleStyleSheet()
        
        # Título
        title = Paragraph(f"Reporte de {tipo_reporte.title()}", styles['Title'])
        elements.append(title)
        elements.append(Spacer(1, 20))
        
        # Período
        period = Paragraph(f"Período: {fecha_inicio} - {fecha_fin}", styles['Normal'])
        elements.append(period)
        elements.append(Spacer(1, 20))
        
        # Obtener datos del reporte
        reportes_view = ReportesViewSet()
        request_mock = type('Request', (), {
            'GET': {'tipo': tipo_reporte, 'fecha_inicio': fecha_inicio, 'fecha_fin': fecha_fin}
        })()
        
        response = reportes_view.get(request_mock)
        data = response.data
        
        # Crear tabla con los datos
        table_data = [['Métrica', 'Valor']]
        for key, value in data.items():
            if key != 'periodo':
                table_data.append([key.replace('_', ' ').title(), str(value)])
        
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 14),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(table)
        doc.build(elements)
        
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="reporte_{tipo_reporte}_{fecha_inicio}_{fecha_fin}.pdf"'
        
        return response
    
    def export_excel(self, tipo_reporte, fecha_inicio, fecha_fin):
        """Exportar reporte en Excel"""
        # Aquí implementarías la exportación a Excel usando openpyxl
        # Por simplicidad, retorno un CSV
        
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="reporte_{tipo_reporte}_{fecha_inicio}_{fecha_fin}.csv"'
        
        writer = csv.writer(response)
        writer.writerow(['Reporte', tipo_reporte.title()])
        writer.writerow(['Período', f'{fecha_inicio} - {fecha_fin}'])
        writer.writerow([])
        writer.writerow(['Métrica', 'Valor'])
        
        # Obtener datos y escribir al CSV
        # ... implementar lógica similar al PDF
        
        return response
